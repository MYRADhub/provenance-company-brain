import shutil

from openpyxl import Workbook

from app import vcs, workbook

DAVID = "david@northbeam.com"
PRIYA = "priya@northbeam.com"


def make_model(path, growth=0.15, margin=0.18, exit_multiple=11.0, note=None):
    """A small LBO-style workbook: labels in column B, years as columns with a header row."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Assumptions"
    ws["B2"], ws["C2"] = "Assumption", "Value"
    ws["B3"], ws["C3"] = "Revenue growth", growth
    ws["B4"], ws["C4"] = "EBITDA margin", margin
    ws["B5"], ws["C5"] = "Exit multiple", exit_multiple
    if note:
        ws["B8"] = note

    om = wb.create_sheet("Operating Model")
    om["B4"] = "($M)"
    years = ["FY25", "FY26", "FY27"]
    for i, y in enumerate(years):
        om.cell(row=4, column=3 + i, value=y)
    revenue = [48.0 * (1 + growth) ** i for i in range(3)]
    om["B5"], om["B6"] = "Revenue ($M)", "EBITDA ($M)"
    for i, r in enumerate(revenue):
        om.cell(row=5, column=3 + i, value=round(r, 4))
        om.cell(row=6, column=3 + i, value=round(r * margin, 4))

    rt = wb.create_sheet("Returns")
    exit_ev = revenue[-1] * margin * exit_multiple
    moic = (exit_ev - 40) / 48
    rt["B3"], rt["C3"] = "Exit EV ($M)", round(exit_ev, 4)
    rt["B4"], rt["C4"] = "Equity IRR", round(moic ** (1 / 5) - 1, 4)
    rt["B5"], rt["C5"] = "MOIC", round(moic, 4)
    wb.save(path)
    return path


def test_read_workbook_labels_use_row_label_and_column_header(tmp_path):
    wb = workbook.read_workbook(make_model(tmp_path / "m.xlsx"))
    om = wb["sheets"]["Operating Model"]
    assert om["D5"]["label"] == "FY26 Revenue ($M)"
    assert om["C6"]["label"] == "FY25 EBITDA ($M)"
    assert wb["sheets"]["Assumptions"]["C3"]["label"] == "Revenue growth"


def test_first_import_creates_model_on_main(tmp_path):
    c = vcs.record_snapshot(make_model(tmp_path / "Falcon LBO v1.xlsx"), "falcon", DAVID, "2026-09-20T10:00:00")
    assert c["branch"] == "main"
    assert c["parent_id"] is None and not c["is_fork"]
    assert c["model"]["name"] == "Falcon LBO"
    assert c["message"] == "Initial import"
    assert c["metrics"]["revenue_growth"] == 0.15
    assert {"fy26_revenue", "ebitda_margin", "irr", "moic"} <= set(c["metrics"])


def test_resave_same_path_adds_child_commit_on_same_branch(tmp_path):
    path = tmp_path / "Falcon LBO.xlsx"
    first = vcs.record_snapshot(make_model(path), "falcon", DAVID, "2026-09-20T10:00:00")
    second = vcs.record_snapshot(make_model(path, growth=0.12), "falcon", PRIYA, "2026-10-02T10:00:00")
    assert second["branch"] == "main"
    assert second["parent_id"] == first["id"]
    assert not second["is_fork"]
    assert second["prior_metrics"]["revenue_growth"] == 0.15
    assert second["metrics"]["revenue_growth"] == 0.12
    growth = next(ch for ch in second["changes"] if ch["sheet"] == "Assumptions" and ch["cell"] == "C3")
    assert (growth["label"], growth["old_value"], growth["new_value"]) == ("Revenue growth", 0.15, 0.12)
    assert "Revenue growth" in second["message"]


def test_identical_content_returns_none(tmp_path):
    path = tmp_path / "Falcon LBO.xlsx"
    vcs.record_snapshot(make_model(path), "falcon", DAVID, "2026-09-20T10:00:00")
    assert vcs.record_snapshot(make_model(path), "falcon", DAVID, "2026-09-21T10:00:00") is None


def test_unedited_copy_returns_none_but_registers_branch(tmp_path):
    src = make_model(tmp_path / "Falcon LBO v1.xlsx")
    first = vcs.record_snapshot(src, "falcon", DAVID, "2026-09-20T10:00:00")
    copy = tmp_path / "Falcon LBO v1 - backup.xlsx"
    shutil.copy(src, copy)
    assert vcs.record_snapshot(copy, "falcon", DAVID, "2026-09-21T10:00:00") is None
    branches = {b["name"]: b["head_commit"] for b in vcs.graph(first["model_id"])["branches"]}
    assert branches["backup"] == first["id"]


def test_versioned_filenames_continue_main_and_copies_fork(tmp_path):
    v1 = vcs.record_snapshot(make_model(tmp_path / "Falcon LBO v1.xlsx"), "falcon", DAVID, "2026-09-20T10:00:00")
    v2 = vcs.record_snapshot(make_model(tmp_path / "Falcon LBO v2.xlsx", growth=0.12), "falcon", PRIYA,
                             "2026-10-02T10:00:00")
    assert v2["branch"] == "main" and v2["parent_id"] == v1["id"] and not v2["is_fork"]

    downside = vcs.record_snapshot(
        make_model(tmp_path / "Falcon LBO v2 - downside case (PS).xlsx", growth=0.06, margin=0.16),
        "falcon", PRIYA, "2026-10-04T10:00:00",
    )
    assert downside["is_fork"]
    assert downside["branch"] == "downside-case-ps"
    assert downside["parent_id"] == v2["id"]
    assert downside["model_id"] == v1["model_id"]

    v3 = vcs.record_snapshot(make_model(tmp_path / "Falcon LBO v3.xlsx", growth=0.08), "falcon", PRIYA,
                             "2026-10-08T10:00:00")
    assert v3["branch"] == "main" and v3["parent_id"] == v2["id"]
    assert vcs.main_head(v1["model_id"]) == v3["id"]


def test_different_structure_creates_separate_model(tmp_path):
    a = vcs.record_snapshot(make_model(tmp_path / "Falcon LBO.xlsx"), "falcon", DAVID, "2026-09-20T10:00:00")
    other = Workbook()
    other.active.title = "Debt Schedule"
    other.active["B2"], other.active["C2"] = "Senior term loan", 60
    other.save(tmp_path / "Falcon debt.xlsx")
    b = vcs.record_snapshot(tmp_path / "Falcon debt.xlsx", "falcon", DAVID, "2026-09-21T10:00:00")
    assert b["model_id"] != a["model_id"] and b["branch"] == "main"
    assert len(vcs.list_models("falcon")) == 2


def test_commit_diff_across_non_adjacent_commits(tmp_path):
    path = tmp_path / "Falcon LBO.xlsx"
    c1 = vcs.record_snapshot(make_model(path), "falcon", DAVID, "2026-09-20T10:00:00")
    vcs.record_snapshot(make_model(path, growth=0.12), "falcon", PRIYA, "2026-10-02T10:00:00")
    c3 = vcs.record_snapshot(make_model(path, growth=0.08, exit_multiple=9.5), "falcon", DAVID, "2026-10-12T10:00:00")
    changes = {(ch["sheet"], ch["cell"]): ch for ch in vcs.commit_diff(c1["id"], c3["id"])}
    assert changes[("Assumptions", "C3")]["old_value"] == 0.15
    assert changes[("Assumptions", "C3")]["new_value"] == 0.08
    assert changes[("Assumptions", "C5")]["new_value"] == 9.5
    assert changes[("Operating Model", "D5")]["label"] == "FY26 Revenue ($M)"
    assert vcs.commit_diff(c1["id"], c1["id"]) == []


def test_stale_references_and_tags(tmp_path):
    path = tmp_path / "Falcon LBO.xlsx"
    c1 = vcs.record_snapshot(make_model(path, growth=0.12), "falcon", PRIYA, "2026-10-02T10:00:00")
    refs = vcs.add_doc_references(
        "falcon_ic_memo_draft", "Project Falcon - IC memo draft", "falcon", DAVID,
        {"revenue_growth": 0.12, "irr": c1["metrics"]["irr"], "exit_multiple": 11.0}, doc_date="2026-10-03",
    )
    assert {r["quoted_commit"] for r in refs} == {c1["id"]}
    assert vcs.stale_references("falcon") == []

    c2 = vcs.record_snapshot(make_model(path, growth=0.08), "falcon", PRIYA, "2026-10-08T10:00:00")
    stale = {s["metric"]: s for s in vcs.stale_references("falcon")}
    assert set(stale) == {"revenue_growth", "irr"}  # exit multiple did not change
    assert stale["revenue_growth"]["latest_value"] == 0.08
    assert stale["revenue_growth"]["latest_commit"] == c2["id"]
    assert stale["revenue_growth"]["quoted_commit"] == c1["id"]

    g = vcs.graph(c1["model_id"])
    tagged = {c["id"]: c["tags"] for c in g["commits"]}
    assert tagged[c1["id"]][0]["doc_id"] == "falcon_ic_memo_draft"
    assert tagged[c2["id"]] == []


def test_graph_shape(tmp_path):
    v1 = vcs.record_snapshot(make_model(tmp_path / "Falcon LBO v1.xlsx"), "falcon", DAVID, "2026-09-20T10:00:00")
    vcs.record_snapshot(make_model(tmp_path / "Falcon LBO v2.xlsx", growth=0.12), "falcon", PRIYA, "2026-10-02T10:00:00")
    fork = vcs.record_snapshot(make_model(tmp_path / "Falcon LBO v2 - downside case (PS).xlsx", growth=0.06),
                               "falcon", PRIYA, "2026-10-04T10:00:00")
    g = vcs.graph(v1["model_id"])
    assert g["model"]["name"] == "Falcon LBO"
    assert [c["committed_at"] for c in g["commits"]] == sorted(c["committed_at"] for c in g["commits"])
    lanes = {b["name"]: b["lane"] for b in g["branches"]}
    assert lanes == {"main": 0, "downside-case-ps": 1}
    fork_node = next(c for c in g["commits"] if c["id"] == fork["id"])
    assert fork_node["lane"] == 1 and fork_node["head_of"] == ["downside-case-ps"] and fork_node["is_head"]
    assert all({"id", "parent_id", "branch", "message", "metrics", "tags"} <= set(c) for c in g["commits"])
    models = vcs.list_models("falcon")
    assert models[0]["commit_count"] == 3 and set(models[0]["branches"]) == {"main", "downside-case-ps"}


def test_resave_of_superseded_file_forks(tmp_path):
    vcs.record_snapshot(make_model(tmp_path / "Falcon LBO v1.xlsx"), "falcon", DAVID, "2026-09-20T10:00:00")
    v2 = vcs.record_snapshot(make_model(tmp_path / "Falcon LBO v2.xlsx", growth=0.12), "falcon", PRIYA,
                             "2026-10-02T10:00:00")
    late = vcs.record_snapshot(make_model(tmp_path / "Falcon LBO v1.xlsx", margin=0.2), "falcon", DAVID,
                               "2026-10-05T10:00:00")
    assert late["is_fork"] and late["branch"] != "main"
    assert vcs.main_head(v2["model_id"]) == v2["id"]


def test_clean_stem_and_branch_suffix():
    assert vcs._clean_stem("Falcon LBO v3 FINAL (2)") == "Falcon LBO"
    assert vcs._branch_suffix("Falcon LBO v3 (DO comments)", "Falcon LBO") == "do-comments"
    assert vcs._branch_suffix("Falcon LBO v4", "Falcon LBO") == ""
