import base64

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from app import governance, ingest, main, store, tokens, vcs, watcher
from tests.test_vcs import make_model

PRIYA = "priya@northbeam.com"
MARCUS = "marcus@northbeam.com"
WB_ID = "wb_0123456789"


@pytest.fixture(autouse=True)
def offline(monkeypatch):
    indexed: list[dict] = []
    monkeypatch.setattr(ingest, "hydra_ingest", lambda items: indexed.extend(items))
    return indexed


@pytest.fixture
def client():
    return TestClient(main.app)


def b64(path) -> str:
    return base64.b64encode(path.read_bytes()).decode()


def save(client, token, path, deal="falcon", **extra):
    body = {"deal": deal, "workbook_id": WB_ID, "file_name": path.name, "data": b64(path), **extra}
    return client.post("/api/v1/saves", json=body, headers={"authorization": f"Bearer {token}"})


def test_me_lists_only_staffed_deals(client):
    res = client.get("/api/v1/me", headers={"authorization": f"Bearer {tokens.create(MARCUS)}"})
    assert res.status_code == 200
    assert res.json()["user"]["name"] == "Marcus Lee"
    assert [d["id"] for d in res.json()["deals"]] == ["harbor"]
    assert res.json()["retrieval"] == "hydra"  # pinned by conftest


def test_requires_a_token(client):
    assert client.get("/api/v1/me").status_code == 401
    assert client.post("/api/v1/saves", json={}, headers={"authorization": "Bearer pb_nope"}).status_code == 401


def test_first_save_starts_a_model_and_the_next_is_a_commit_on_it(client, tmp_path, offline):
    token = tokens.create(PRIYA)
    path = make_model(tmp_path / "Falcon LBO.xlsx")
    first = save(client, token, path, saved_at=1_790_000_000_000)
    assert first.status_code == 200
    c1 = first.json()["commit"]
    assert c1["branch"] == "main" and c1["parent_id"] is None
    assert c1["model"]["name"] == "Falcon LBO"

    make_model(path, growth=0.20)
    c2 = save(client, token, path, saved_at=1_790_000_600_000).json()["commit"]
    assert c2["parent_id"] == c1["id"]
    assert c2["branch"] == "main"
    assert c2["changed_cells"] > 0
    assert c2["metrics"]["revenue_growth"] == 0.20 and c2["prior_metrics"]["revenue_growth"] == 0.15

    commit = vcs.get_commit(c2["id"])
    assert commit["author_email"] == PRIYA
    # The change log is indexed and registered so the brain can find it and link it.
    assert any(d.get("id") == f"f_commit_{c2['id']}" or f"f_commit_{c2['id']}" in str(d) for d in offline)
    assert store.get_document(f"f_commit_{c2['id']}")
    # Proactive alerts pick it up like a folder save.
    assert any(ev.get("commit", {}).get("id") == c2["id"] for ev in watcher.events_since(0))


def test_saving_the_same_content_again_records_nothing(client, tmp_path):
    token = tokens.create(PRIYA)
    path = make_model(tmp_path / "Falcon LBO.xlsx")
    save(client, token, path)
    again = save(client, token, path)
    assert again.status_code == 200 and again.json()["commit"] is None


def test_walled_deal_is_not_found_and_audited(client, tmp_path):
    token = tokens.create(MARCUS)
    walled = save(client, token, make_model(tmp_path / "x.xlsx"), deal="falcon")
    unknown = save(client, token, make_model(tmp_path / "y.xlsx"), deal="nope")
    assert walled.status_code == unknown.status_code == 404
    assert walled.json() == unknown.json()
    assert (MARCUS, "blocked") in [(e["email"], e["action"]) for e in governance.audit_log()]


def test_reports_documents_that_now_quote_old_numbers(client, tmp_path):
    token = tokens.create(PRIYA)
    path = make_model(tmp_path / "Falcon LBO.xlsx")
    c1 = save(client, token, path).json()["commit"]
    store.execute(
        "INSERT INTO doc_references VALUES (?,?,?,?,?,?,?,?)",
        ("memo", "IC memo", "falcon", c1["model"]["id"], "revenue_growth", 0.15, c1["id"], "david@northbeam.com"),
    )
    make_model(path, growth=0.25)
    res = save(client, token, path).json()
    assert res["stale"] == [{"document": "IC memo", "metric": "revenue_growth", "quoted_value": 0.15,
                             "latest_value": 0.25, "owner": "David Okafor"}]


def test_tessel_metadata_sheet_is_not_part_of_the_model(client, tmp_path):
    token = tokens.create(PRIYA)
    path = make_model(tmp_path / "Falcon LBO.xlsx")
    save(client, token, path)
    wb = load_workbook(path)
    meta = wb.create_sheet("_tessel_meta")
    meta["A1"] = '{"aiTurns": []}'
    meta.sheet_state = "veryHidden"
    wb.save(path)
    assert save(client, token, path).json()["commit"] is None


@pytest.mark.parametrize(
    "patch,status",
    [
        ({"workbook_id": "../../etc"}, 400),
        ({"file_name": "notes.txt"}, 400),
        ({"data": "not base64!"}, 400),
        ({"data": base64.b64encode(b"hello").decode()}, 400),
    ],
)
def test_rejects_bad_input(client, tmp_path, patch, status):
    token = tokens.create(PRIYA)
    path = make_model(tmp_path / "Falcon LBO.xlsx")
    body = {"deal": "falcon", "workbook_id": WB_ID, "file_name": path.name, "data": b64(path), **patch}
    res = client.post("/api/v1/saves", json=body, headers={"authorization": f"Bearer {token}"})
    assert res.status_code == status


def test_file_name_cannot_escape_the_upload_folder(client, tmp_path):
    token = tokens.create(PRIYA)
    path = make_model(tmp_path / "Falcon LBO.xlsx")
    body = {"deal": "falcon", "workbook_id": WB_ID, "file_name": "../../../evil.xlsx", "data": b64(path)}
    res = client.post("/api/v1/saves", json=body, headers={"authorization": f"Bearer {token}"})
    assert res.status_code == 200
    stored = store._rows("SELECT file_path FROM commits")[0]["file_path"]
    assert stored.endswith(f"uploads/falcon/{WB_ID}/evil.xlsx")


def test_unreadable_workbook_is_reported(client):
    token = tokens.create(PRIYA)
    body = {"deal": "falcon", "workbook_id": WB_ID, "file_name": "x.xlsx",
            "data": base64.b64encode(b"PK\x03\x04 broken").decode()}
    res = client.post("/api/v1/saves", json=body, headers={"authorization": f"Bearer {token}"})
    assert res.status_code == 422


def test_browser_origin_rules(client):
    token = tokens.create(PRIYA)
    ok = client.get("/api/v1/me", headers={"authorization": f"Bearer {token}", "origin": "http://localhost:5173"})
    bad = client.get("/api/v1/me", headers={"authorization": f"Bearer {token}", "origin": "https://evil.example"})
    assert ok.status_code == 200 and ok.headers["access-control-allow-origin"] == "http://localhost:5173"
    assert bad.status_code == 403
    preflight = client.options("/api/v1/saves", headers={
        "origin": "http://localhost:5173", "access-control-request-method": "POST",
        "access-control-request-headers": "authorization,content-type"})
    assert preflight.status_code == 200


def history(client, token, **params):
    return client.get("/api/v1/history", params={"deal": "falcon", "workbook_id": WB_ID, **params},
                      headers={"authorization": f"Bearer {token}"})


def test_history_lists_versions_and_one_cells_changes(client, tmp_path):
    token = tokens.create(PRIYA)
    path = make_model(tmp_path / "Falcon LBO.xlsx")
    save(client, token, path, saved_at=1_790_000_000_000)
    make_model(path, growth=0.15, margin=0.2)  # growth unchanged in this version
    save(client, token, path, saved_at=1_790_000_600_000)
    make_model(path, growth=0.2, margin=0.2)
    save(client, token, path, saved_at=1_790_001_200_000)

    res = history(client, token, sheet="Assumptions", cell="C3").json()
    assert len(res["versions"]) == 3
    assert res["versions"][0]["author"] == "Priya Shah"
    assert res["versions"][0]["committed_at"] > res["versions"][-1]["committed_at"]
    assert res["cell"]["label"] == "Revenue growth"
    assert [c["value"] for c in res["cell"]["changes"]] == [0.15, 0.2]
    assert res["cell"]["changes"][-1]["commit"] == res["versions"][0]["commit"]


def test_history_of_an_unsaved_workbook_is_empty(client):
    res = history(client, tokens.create(PRIYA), sheet="Assumptions", cell="C3").json()
    assert res == {"versions": [], "cell": {"sheet": "Assumptions", "cell": "C3", "label": "", "changes": []}}


def test_history_respects_walls(client, tmp_path):
    save(client, tokens.create(PRIYA), make_model(tmp_path / "Falcon LBO.xlsx"))
    res = history(client, tokens.create(MARCUS))
    assert res.status_code == 404


@pytest.mark.parametrize("params", [{"sheet": "Assumptions"}, {"sheet": "A", "cell": "c3"}, {"sheet": "A", "cell": "A0"},
                                    {"workbook_id": "../x"}])
def test_history_rejects_bad_queries(client, params):
    assert history(client, tokens.create(PRIYA), **params).status_code == 400
