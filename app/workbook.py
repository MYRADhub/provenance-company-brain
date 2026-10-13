"""Read Excel workbooks into a normalized cell map, hash them, and diff them.

A workbook is represented as {"sheets": {sheet: {cell: {"value", "formula", "label"}}}, "order": [sheets]}.
`value` is the cached value (what Excel last computed), `formula` the formula text if the cell has one.
`label` is a human-readable name for the cell built from the nearest text to its left in the same row,
prefixed by the column header above when the row holds several values (e.g. "FY26 Revenue ($M)").
"""

import datetime as dt
import hashlib
import json
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

HEADER_SCAN_ROWS = 40
# Tessel keeps its own metadata (provenance, AI log) in a very hidden sheet; it is not part of the model.
APP_SHEETS = {"_tessel_meta"}


def _norm(value):
    if isinstance(value, bool):
        return value
    if isinstance(value, float):
        value = round(value, 6)
        return int(value) if value.is_integer() else value
    if isinstance(value, (dt.datetime, dt.date, dt.time)):
        return value.isoformat()
    if isinstance(value, str):
        value = value.strip()
        return value or None
    return value


def _is_text(value) -> bool:
    return isinstance(value, str) and bool(value.strip())


def read_workbook(path: str | Path) -> dict:
    """Load both the formulas and the cached values of every non-empty cell."""
    wb_f = load_workbook(path, data_only=False)
    wb_v = load_workbook(path, data_only=True)
    sheets: dict[str, dict] = {}
    names = [n for n in wb_f.sheetnames if not (n in APP_SHEETS and wb_f[n].sheet_state == "veryHidden")]
    for name in names:
        ws_f = wb_f[name]
        ws_v = wb_v[name]
        grid: dict[tuple[int, int], dict] = {}
        for row in ws_f.iter_rows():
            for c in row:
                raw = c.value
                if raw is None:
                    continue
                formula = raw if isinstance(raw, str) and raw.startswith("=") else None
                value = _norm(ws_v.cell(row=c.row, column=c.column).value) if formula else _norm(raw)
                if value is None and formula is None:
                    continue
                grid[(c.row, c.column)] = {"value": value, "formula": formula}
        sheets[name] = _label_grid(grid)
    return {"order": names, "sheets": sheets}


def _label_grid(grid: dict[tuple[int, int], dict]) -> dict[str, dict]:
    values_per_row: dict[int, int] = {}
    for (r, _c), cell in grid.items():
        if not _is_text(cell["value"]):
            values_per_row[r] = values_per_row.get(r, 0) + 1

    out: dict[str, dict] = {}
    for (r, c), cell in sorted(grid.items()):
        row_label = None
        for cc in range(c - 1, 0, -1):
            left = grid.get((r, cc))
            if left and _is_text(left["value"]):
                row_label = left["value"]
                break
        header = None
        if values_per_row.get(r, 0) >= 2 and not _is_text(cell["value"]):
            for rr in range(r - 1, max(0, r - HEADER_SCAN_ROWS), -1):
                above = grid.get((rr, c))
                if above is None:
                    continue
                v = above["value"]
                if _is_text(v) or (isinstance(v, int) and 1900 <= v <= 2100):
                    header = str(v)
                    break
        if row_label and header:
            label = f"{header} {row_label}"
        else:
            label = row_label or header or ""
        out[f"{get_column_letter(c)}{r}"] = {**cell, "label": label}
    return out


def cell_map(wb: dict) -> dict[str, object]:
    """Flat {sheet!cell: formula or value}; the identity of a workbook's content."""
    flat: dict[str, object] = {}
    for sheet, cells in wb["sheets"].items():
        for ref, cell in cells.items():
            flat[f"{sheet}!{ref}"] = cell["formula"] if cell["formula"] is not None else cell["value"]
    return flat


def cell_map_hash(wb: dict) -> str:
    payload = json.dumps({"order": wb["order"], "cells": cell_map(wb)}, sort_keys=True, default=str)
    return hashlib.sha256(payload.encode()).hexdigest()


def similarity(wb_a: dict, wb_b: dict) -> float:
    """Fraction of identical cells over the union of non-empty cells."""
    a, b = cell_map(wb_a), cell_map(wb_b)
    keys = set(a) | set(b)
    if not keys:
        return 1.0
    same = sum(1 for k in keys if k in a and k in b and a[k] == b[k])
    return same / len(keys)


def structure(wb: dict) -> set[str]:
    """Sheet names plus row labels: what makes two files 'the same model'."""
    sig = {f"sheet:{s}" for s in wb["order"]}
    for sheet, cells in wb["sheets"].items():
        for cell in cells.values():
            if _is_text(cell["value"]):
                sig.add(f"{sheet}:{cell['value'].lower()}")
    return sig


def structure_similarity(wb_a: dict, wb_b: dict) -> float:
    a, b = structure(wb_a), structure(wb_b)
    if not a and not b:
        return 1.0
    return len(a & b) / len(a | b)


def _sort_key(order: list[str], sheet: str, ref: str):
    col = "".join(ch for ch in ref if ch.isalpha())
    row = int("".join(ch for ch in ref if ch.isdigit()) or 0)
    idx = order.index(sheet) if sheet in order else len(order)
    return (idx, row, len(col), col)


def diff(wb_a: dict | None, wb_b: dict) -> list[dict]:
    """Cell-level changes from wb_a to wb_b (value or formula differs)."""
    wb_a = wb_a or {"order": [], "sheets": {}}
    changes = []
    sheets = list(dict.fromkeys(wb_b["order"] + wb_a["order"]))
    for sheet in sheets:
        a = wb_a["sheets"].get(sheet, {})
        b = wb_b["sheets"].get(sheet, {})
        for ref in set(a) | set(b):
            old, new = a.get(ref), b.get(ref)
            old_v, new_v = (old or {}).get("value"), (new or {}).get("value")
            old_f, new_f = (old or {}).get("formula"), (new or {}).get("formula")
            if old_v == new_v and old_f == new_f:
                continue
            changes.append({
                "sheet": sheet,
                "cell": ref,
                "label": (new or {}).get("label") or (old or {}).get("label") or "",
                "old_value": old_v,
                "new_value": new_v,
                "old_formula": old_f,
                "new_formula": new_f,
            })
    order = sheets
    changes.sort(key=lambda ch: _sort_key(order, ch["sheet"], ch["cell"]))
    return changes


def dump_for_llm(wb: dict, max_cells: int = 600) -> str:
    """Compact text listing of labelled numeric cells, for metric discovery."""
    lines = []
    for sheet in wb["order"]:
        for ref, cell in wb["sheets"][sheet].items():
            v = cell["value"]
            if isinstance(v, (int, float)) and not isinstance(v, bool) and cell["label"]:
                lines.append(f"{sheet}!{ref} | {cell['label']} | {v}")
                if len(lines) >= max_cells:
                    return "\n".join(lines)
    return "\n".join(lines)
