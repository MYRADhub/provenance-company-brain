"""Cell-level diff between two workbooks."""
from pathlib import Path

from openpyxl import load_workbook


def _norm(value):
    if isinstance(value, float):
        return round(value, 6)
    return value


def diff_workbooks(path_a: str | Path, path_b: str | Path) -> list[dict]:
    """Return every cell whose value differs between two workbooks.

    Each change is {sheet, cell, label, old_value, new_value}; label is the text in
    column A of the same row, so a change reads as "FY26 Revenue: 53.76 -> 51.84".
    """
    wb_a = load_workbook(path_a, data_only=True)
    wb_b = load_workbook(path_b, data_only=True)
    changes: list[dict] = []
    for sheet in wb_b.sheetnames:
        ws_b = wb_b[sheet]
        ws_a = wb_a[sheet] if sheet in wb_a.sheetnames else None
        max_row = max(ws_b.max_row, ws_a.max_row if ws_a else 0)
        max_col = max(ws_b.max_column, ws_a.max_column if ws_a else 0)
        for row in range(1, max_row + 1):
            label_b = ws_b.cell(row=row, column=1).value
            label_a = ws_a.cell(row=row, column=1).value if ws_a else None
            label = label_b if label_b is not None else label_a
            for col in range(1, max_col + 1):
                new = ws_b.cell(row=row, column=col).value
                old = ws_a.cell(row=row, column=col).value if ws_a else None
                if _norm(old) == _norm(new):
                    continue
                changes.append({
                    "sheet": sheet,
                    "cell": ws_b.cell(row=row, column=col).coordinate,
                    "label": str(label) if label is not None else "",
                    "old_value": old,
                    "new_value": new,
                })
    for sheet in wb_a.sheetnames:
        if sheet not in wb_b.sheetnames:
            changes.append({"sheet": sheet, "cell": "", "label": "sheet removed",
                            "old_value": sheet, "new_value": None})
    return changes
