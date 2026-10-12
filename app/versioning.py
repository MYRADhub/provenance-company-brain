"""Model versioning: write the workbook, diff it against the prior version, record it."""
import json
import time
from pathlib import Path

from openpyxl import Workbook

from app import store
from app.config import COLLECTION, HYDRA_DATABASE, MODELS_DIR, hydra
from app.excel_diff import diff_workbooks
from app.seed_content import (
    BASE_FY25_REVENUE, DEALS, ENTRY_EQUITY, ENTRY_EV, EXIT_NET_DEBT, HOLD_YEARS,
    USERS, compute_model, deal_acl, version_spec,
)

PERCENT_LABELS = {"Revenue Growth (annual)", "EBITDA Margin", "IRR"}
METRIC_LABELS = {
    "revenue_growth": "Revenue growth", "ebitda_margin": "EBITDA margin", "fy26_revenue": "FY26 revenue",
    "fy26_ebitda": "FY26 EBITDA", "exit_multiple": "Exit multiple", "irr": "Gross IRR", "moic": "MOIC",
}


def _user_name(email: str) -> str:
    return next((u["name"] for u in USERS if u["email"] == email), email)


def _deal(deal_id: str) -> dict:
    return next(d for d in DEALS if d["id"] == deal_id)


def workbook_path(deal_id: str, version: int) -> Path:
    return MODELS_DIR / f"{deal_id}_model_v{version}.xlsx"


def build_workbook(deal_id: str, version: int) -> Path:
    spec = version_spec(deal_id, version)
    model = compute_model(spec["growth"], spec["margin"], spec["exit_multiple"])
    wb = Workbook()
    ws = wb.active
    ws.title = "Assumptions"
    rows = [
        ("Assumption", "Value"),
        ("Base FY25 Revenue ($M)", BASE_FY25_REVENUE),
        ("Revenue Growth (annual)", spec["growth"]),
        ("EBITDA Margin", spec["margin"]),
        ("Entry Enterprise Value ($M)", ENTRY_EV),
        ("Entry Equity ($M)", ENTRY_EQUITY),
        ("Net Debt at Exit ($M)", EXIT_NET_DEBT),
        ("Exit Multiple (x FY30 EBITDA)", spec["exit_multiple"]),
        ("Hold Period (years)", HOLD_YEARS),
    ]
    for r in rows:
        ws.append(r)
    rs = wb.create_sheet("Returns")
    rs.append(("Output", "Value"))
    for year, value in model["revenue"].items():
        rs.append((f"FY{year % 100} Revenue ($M)", value))
    for year, value in model["ebitda"].items():
        rs.append((f"FY{year % 100} EBITDA ($M)", value))
    rs.append(("Exit Enterprise Value ($M)", model["exit_ev"]))
    rs.append(("Exit Equity Value ($M)", model["exit_equity"]))
    rs.append(("MOIC", model["metrics"]["moic"]))
    rs.append(("IRR", model["metrics"]["irr"]))
    for sheet in (ws, rs):
        sheet.column_dimensions["A"].width = 32
        sheet.column_dimensions["B"].width = 14
        for row in sheet.iter_rows(min_row=2):
            label, cell = row[0].value, row[1]
            cell.number_format = "0.0%" if label in PERCENT_LABELS else "0.00"
    MODELS_DIR.mkdir(parents=True, exist_ok=True)
    path = workbook_path(deal_id, version)
    wb.save(path)
    return path


def _fmt(label: str, value) -> str:
    if isinstance(value, (int, float)) and label in PERCENT_LABELS:
        return f"{value * 100:.1f}%"
    if isinstance(value, float):
        return f"{value:.2f}"
    return str(value)


def _fmt_metric(metric: str, value) -> str:
    if value is None:
        return "n/a"
    if metric in ("revenue_growth", "ebitda_margin", "irr"):
        return f"{value * 100:.1f}%"
    if metric in ("exit_multiple", "moic"):
        return f"{value:.2f}x" if metric == "moic" else f"{value:.1f}x"
    return f"${value:.2f}M"


def changelog_doc(deal_id: str, version: int, author_email: str, saved_at: str,
                  changes: list[dict], metrics: dict, prior_metrics: dict) -> dict:
    deal = _deal(deal_id)
    date = saved_at[:10]
    if version == 1:
        lines = [f"{deal['codename']} operating model v1 created by {_user_name(author_email)} on {date}.",
                 "Initial version. Key outputs:"]
        lines += [f"- {METRIC_LABELS[m]}: {_fmt_metric(m, v)}" for m, v in metrics.items()]
    else:
        lines = [f"{deal['codename']} operating model change log: v{version - 1} -> v{version}.",
                 f"Saved by {_user_name(author_email)} on {date} from Excel. {len(changes)} cells changed.",
                 "", "Cells changed:"]
        lines += [f"- {c['sheet']}!{c['cell']} {c['label']}: {_fmt(c['label'], c['old_value'])} -> "
                  f"{_fmt(c['label'], c['new_value'])}" for c in changes]
        lines += ["", "Key outputs:"]
        lines += [f"- {METRIC_LABELS[m]}: {_fmt_metric(m, prior_metrics.get(m))} -> {_fmt_metric(m, v)}"
                  for m, v in metrics.items() if prior_metrics.get(m) != v]
        lines += ["", "This log is generated from the workbook diff. It records what changed, not why."]
    return {
        "doc_id": f"{deal_id}_model_changelog_v{version}",
        "title": f"{deal['codename']} - model change log v{version}",
        "deal_id": deal_id, "kind": "knowledge_base", "doc_kind": "changelog",
        "date": date, "author_email": author_email, "acl": deal_acl(deal_id),
        "body": "\n".join(lines),
    }


# --------------------------------------------------------------------- HydraDB

def to_app_item(doc: dict) -> dict:
    kind = doc["kind"]
    ts = f"{doc['date']}T12:00:00Z"
    author = _user_name(doc["author_email"])
    fields: dict = {"kind": kind, "body": doc["body"]}
    if kind == "email":
        fields.update({"subject": doc["title"], "from": doc.get("from", doc["author_email"]),
                       "to": doc.get("to", []), "cc": doc.get("cc", []), "created_at": ts})
    elif kind == "meeting":
        fields.update({"title": doc["title"], "host": author, "host_email": doc["author_email"],
                       "participants": [{"name": p} for p in doc.get("participants", [])], "started_at": ts})
    elif kind == "message":
        fields.update({"author": author, "created_at": ts})
    else:
        fields.update({"title": doc["title"], "created_by": author, "created_at": ts})
    return {
        "id": doc["doc_id"],
        "database": HYDRA_DATABASE,
        "collection": COLLECTION,
        "title": doc["title"],
        "kind": kind,
        "provider": "northbeam",
        "external_id": doc["doc_id"],
        "timestamp": ts,
        "fields": fields,
        "metadata": {"deal": doc["deal_id"] or "firm", "doc_kind": doc["doc_kind"]},
        "additional_metadata": {"author": doc["author_email"], "date": doc["date"]},
        "acl": doc["acl"],
    }


def record_document(doc: dict) -> None:
    store.execute(
        "INSERT OR REPLACE INTO documents(doc_id, title, deal_id, kind, date, author_email, acl_json) "
        "VALUES (?,?,?,?,?,?,?)",
        (doc["doc_id"], doc["title"], doc["deal_id"], doc["doc_kind"], doc["date"], doc["author_email"],
         json.dumps(doc["acl"])),
    )


def ingest_documents(docs: list[dict], wait: bool = True, batch_size: int = 10, timeout: float = 300) -> list[str]:
    client = hydra()
    ids: list[str] = []
    for i in range(0, len(docs), batch_size):
        batch = docs[i:i + batch_size]
        client.context.ingest(
            type="knowledge", database=HYDRA_DATABASE, collection=COLLECTION, upsert="true",
            app_knowledge=json.dumps([to_app_item(d) for d in batch]),
        )
        ids += [d["doc_id"] for d in batch]
    for d in docs:
        record_document(d)
    if wait:
        wait_until_searchable(ids, timeout=timeout)
    return ids


def wait_until_searchable(ids: list[str], timeout: float = 300) -> dict:
    client = hydra()
    deadline = time.time() + timeout
    pending = set(ids)
    statuses: dict = {}
    while pending and time.time() < deadline:
        resp = client.context.status(database=HYDRA_DATABASE, collection=COLLECTION, ids=sorted(pending))
        for s in resp.data.statuses:
            statuses[s.id] = s.indexing_status
            if s.indexing_status in ("graph_creation", "completed"):
                pending.discard(s.id)
            elif s.indexing_status in ("errored", "failed"):
                raise RuntimeError(f"indexing failed for {s.id}: {s.error_code} {s.error_message or s.message}")
        if pending:
            time.sleep(3)
    if pending:
        raise TimeoutError(f"still indexing after {timeout}s: {sorted(pending)}")
    return statuses


# ------------------------------------------------------------------ save flow

def save_version(deal_id: str, version: int | None = None, wait: bool = False) -> dict:
    """Write the workbook for a version, diff it against the prior one, and record everything."""
    if version is None:
        version = (store.latest_version(deal_id) or 0) + 1
    spec = version_spec(deal_id, version)
    path = build_workbook(deal_id, version)
    metrics = compute_model(spec["growth"], spec["margin"], spec["exit_multiple"])["metrics"]
    prior_metrics: dict = {}
    changes: list[dict] = []
    if version > 1:
        prior = workbook_path(deal_id, version - 1)
        if not prior.exists():
            build_workbook(deal_id, version - 1)
        changes = diff_workbooks(prior, path)
        prior_metrics = store.get_metrics(deal_id, version - 1)

    store.execute("DELETE FROM model_versions WHERE deal_id = ? AND version = ?", (deal_id, version))
    store.execute("DELETE FROM cell_changes WHERE deal_id = ? AND to_version = ?", (deal_id, version))
    store.execute("DELETE FROM key_metrics WHERE deal_id = ? AND version = ?", (deal_id, version))
    store.execute(
        "INSERT INTO model_versions(deal_id, version, file_path, author_email, saved_at, note) VALUES (?,?,?,?,?,?)",
        (deal_id, version, str(path), spec["author_email"], spec["saved_at"], spec["note"]),
    )
    store.executemany(
        "INSERT INTO cell_changes(deal_id, from_version, to_version, sheet, cell, label, old_value, new_value) "
        "VALUES (?,?,?,?,?,?,?,?)",
        [(deal_id, version - 1, version, c["sheet"], c["cell"], c["label"],
          json.dumps(c["old_value"]), json.dumps(c["new_value"])) for c in changes],
    )
    store.executemany(
        "INSERT INTO key_metrics(deal_id, version, metric, value) VALUES (?,?,?,?)",
        [(deal_id, version, m, v) for m, v in metrics.items()],
    )
    doc = changelog_doc(deal_id, version, spec["author_email"], spec["saved_at"], changes, metrics, prior_metrics)
    ingest_documents([doc], wait=wait)
    return {
        "version": version,
        "author": spec["author_email"],
        "saved_at": spec["saved_at"],
        "changes": changes,
        "metrics": metrics,
        "prior_metrics": prior_metrics,
        "changelog_doc_id": doc["doc_id"],
    }
